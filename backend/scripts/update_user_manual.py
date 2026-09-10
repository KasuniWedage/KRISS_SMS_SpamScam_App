from pathlib import Path

from docx import Document


SOURCE = Path(
    r"C:\Users\Admin\Downloads\Kriss App\KRISS_SMS_Shield_User_Manual_Real_Data.docx"
)
OUTPUT = Path(__file__).resolve().parents[2] / "docs" / "KRISS_SMS_Shield_User_Manual_Updated.docx"


def replace_paragraph(document: Document, starts_with: str, replacement: str) -> None:
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(starts_with):
            paragraph.text = replacement
            return
    raise ValueError(f"Paragraph not found: {starts_with}")


def add_before(anchor, text: str, style: str | None = None):
    paragraph = anchor.insert_paragraph_before(text)
    if style:
        paragraph.style = style
    return paragraph


def main() -> None:
    document = Document(SOURCE)

    # Correct product, authentication, model, and evidence descriptions.
    replace_paragraph(
        document,
        "This manual covers the KRISS SMS Shield Android client",
        "This manual covers the KRISS SMS Shield Android client and the FastAPI "
        "backend functions used by end users and administrators. Model training is "
        "documented separately in colab/KRISS_SMS_Spam_Scam_Training_V2.ipynb and "
        "is summarized in Appendix B.",
    )
    replace_paragraph(
        document,
        "User authentication (email/password, Google, and Facebook)",
        "Email/password accounts are authenticated by the FastAPI backend. Google "
        "and Facebook sign-in use Firebase Authentication on Android; the Firebase "
        "identity token is verified by the backend before an application JWT session "
        "is issued.",
    )
    replace_paragraph(
        document,
        "Note: The /health response should report",
        "Note: The /health response should report model_ready = true and model version "
        "KRISS-Colab-2.0-Trilingual. If model_ready is false, confirm that "
        "backend/model/sms_classifier_v2.joblib and "
        "backend/model/model_metadata_v2.json are present and that installed package "
        "versions match backend/model/model_runtime_requirements.txt.",
    )
    replace_paragraph(
        document,
        "Registration is processed through Firebase Authentication.",
        "Email/password registration is processed by the backend. Passwords are "
        "stored only as secure hashes. Google and Facebook accounts are processed "
        "through Firebase and synchronized with a backend profile after sign-in.",
    )
    replace_paragraph(
        document,
        "Passwords for email/password accounts are managed by Firebase Authentication",
        "Passwords for email/password accounts are hashed by the backend and are "
        "never stored as plain text. Google and Facebook credentials remain managed "
        "by their identity providers through Firebase Authentication.",
    )
    replace_paragraph(
        document,
        "The backend records an audit log",
        "The backend records key account, administration, and message actions in an "
        "append-only SHA-256 hash chain. Verification can detect missing or modified "
        "records. This design is cryptographically tamper-evident; production "
        "deployments should store exported chain-head anchors in a separate immutable "
        "or append-only system.",
    )
    replace_paragraph(
        document,
        "Yes. The system detects English, Sinhala, and Singlish.",
        "Yes. The active model is trained for English, Sinhala, and Singlish. Tamil "
        "is outside the active model scope and should not be presented as a supported "
        "classification language.",
    )
    replace_paragraph(
        document,
        "The bundled 10,000-message training dataset",
        "The active trilingual model was trained and evaluated using a 10,000-message "
        "localized development dataset. Template-generated, localized, and consented "
        "records must remain distinguishable, and any research or production claim "
        "must be supported by documented source and consent provenance.",
    )
    replace_paragraph(
        document,
        "Tamil-language messages are not validated",
        "Tamil-language classification is not part of the active model scope. The "
        "supported model languages are English, Sinhala, and Singlish.",
    )
    replace_paragraph(
        document,
        "End-to-end response time depends",
        "End-to-end response time depends on the deployed backend, device, and network. "
        "Local API benchmarks verify server-side integration performance but do not "
        "represent physical-device or production-network latency.",
    )
    replace_paragraph(
        document,
        "colab/KRISS_SMS_Spam_Scam_Training.ipynb",
        "colab/KRISS_SMS_Spam_Scam_Training_V2.ipynb — Google Colab notebook "
        "documenting preprocessing, model comparison, training, and evaluation.",
    )
    replace_paragraph(
        document,
        "Active model: KRISS-Colab-2.0,",
        "Active model: KRISS-Colab-2.0-Trilingual, a complete scikit-learn pipeline "
        "combining word-level and character-level TF-IDF features with a Multinomial "
        "Naive Bayes classifier (alpha = 0.35).",
    )
    replace_paragraph(
        document,
        "Bundled development dataset:",
        "Active model dataset: 10,000 localized development messages (4,000 "
        "Legitimate, 3,000 Spam, and 3,000 Scam) in English, Sinhala, and Singlish.",
    )
    replace_paragraph(
        document,
        "Reported metadata results:",
        "Reported metadata results: 0.883 test accuracy and approximately 0.8322 "
        "scam recall. The required minimum accuracy is 0.90; 0.95 is an aspirational "
        "target. The restored model therefore remains below the stated minimum.",
    )
    replace_paragraph(
        document,
        "Important: The reported results describe",
        "Important: These results describe the notebook evaluation split and are not "
        "a guarantee of performance on unseen messages. Final research claims require "
        "an independent English/Sinhala/Singlish test set with documented provenance.",
    )

    # Add current administrator and operations guidance before troubleshooting.
    anchor = next(p for p in document.paragraphs if p.text.strip().startswith("10. Error Messages"))
    add_before(anchor, "9A. Administrator and Dataset Management", "Heading 1")
    add_before(anchor, "9A.1 Opening the Administrator Console", "Heading 2")
    add_before(
        anchor,
        "Accounts with the admin role are directed to the Administrator screen after "
        "sign-in. Administrators can also open the backend web console at /admin. "
        "Normal user accounts cannot access administrator API operations.",
    )
    for text in (
        "Review system totals, classification distribution, API latency, and centralized error records.",
        "View registered users and update user roles with appropriate authorization.",
        "Create database backups, list available backups, and execute isolated restoration tests.",
        "Verify the cryptographically tamper-evident audit chain.",
        "Open Dataset Management to review pending samples and approve or reject governed records.",
    ):
        add_before(anchor, text, "List Paragraph")
    add_before(anchor, "9A.2 Dataset Governance", "Heading 2")
    add_before(
        anchor,
        "A dataset submission includes the message, expected label, language, category, "
        "source type, source reference, and consent reference where applicable. "
        "Administrators must verify provenance before approval. Approved records can "
        "be exported as CSV; approval does not automatically retrain or replace the "
        "active model.",
    )
    add_before(anchor, "9A.3 Backup and Monitoring Operations", "Heading 2")
    add_before(
        anchor,
        "The presence of backup scripts does not activate daily backups. An administrator "
        "must run backend/scripts/setup_backup_scheduler.ps1 and confirm the Windows "
        "scheduled task and backup history. Uptime reports require genuine health-probe "
        "logs and must not be generated from invented observations.",
    )

    # Update cover metadata and feature/technology tables.
    document.tables[0].cell(1, 1).text = "1.0 (Active ML model: KRISS-Colab-2.0-Trilingual)"
    document.tables[0].cell(2, 1).text = "1.1"
    document.tables[0].cell(3, 1).text = "31 August 2026"
    document.tables[0].cell(6, 1).text = "Updated"
    document.tables[1].cell(5, 1).text = (
        "scikit-learn: word and character TF-IDF with a Multinomial Naive Bayes "
        "classifier (English, Sinhala, and Singlish)"
    )
    document.tables[2].cell(9, 0).text = "Administration"
    document.tables[2].cell(9, 1).text = (
        "Metrics, user-role management, error monitoring, backup and restoration "
        "testing, audit verification, and dataset governance"
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
