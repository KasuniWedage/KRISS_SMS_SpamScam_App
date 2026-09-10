from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Admin\Downloads\KRISS_SMS_Shield_User_Manual (1).docx")
OUTPUT = Path(__file__).resolve().parents[2] / "docs" / "KRISS_SMS_Shield_User_Manual_v2.1_Corrected.docx"


def replace_paragraph(document: Document, prefix: str, text: str) -> None:
    matches = [p for p in document.paragraphs if p.text.strip().startswith(prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one paragraph for {prefix!r}; found {len(matches)}")
    matches[0].text = text


def main() -> None:
    document = Document(SOURCE)

    replacements = {
        "This manual covers the KRISS SMS Shield Android client":
            "This manual covers the KRISS SMS Shield Android client and the FastAPI "
            "backend functions used by end users and administrators. Internal model "
            "training is documented separately in "
            "colab/KRISS_SMS_Spam_Scam_Training_V2.ipynb and is summarized in Appendix B.",
        "User authentication (email/password, Google, and Facebook)":
            "Email/password accounts are registered and authenticated by the FastAPI "
            "backend. Google and Facebook sign-in use Firebase Authentication on "
            "Android; the resulting Firebase identity token is verified by the backend "
            "before an application JWT session is issued.",
        "Note: The /health response should report":
            "Note: The /health response should report model_ready = true and model "
            "version KRISS-Colab-2.0-Trilingual. If model_ready is false, confirm that "
            "backend/model/sms_classifier_v2.joblib and "
            "backend/model/model_metadata_v2.json are present and that installed "
            "package versions match backend/model/model_runtime_requirements.txt.",
        "Registration is processed through Firebase Authentication.":
            "Email/password registration is processed by the backend, which stores only "
            "a password hash. Google and Facebook identities are handled through "
            "Firebase and synchronized with a backend user profile after sign-in.",
        "Every completed batch run is saved automatically to your account":
            "Every completed batch run is saved to the authenticated account on the "
            "server and cached in encrypted local storage for responsive display. The "
            "screen synchronizes with server history when the backend is reachable.",
        "Choose the report language: English, Sinhala, or Tamil.":
            "Choose the available report language shown by the current screen. Single "
            "and batch report screens offer English, Sinhala, and Tamil; history-item "
            "reports currently offer English and Sinhala.",
        "Language – English, Sinhala, or Tamil.":
            "Interface language – English, Sinhala, or Tamil. This setting localizes "
            "the application interface; it does not expand the active model's supported "
            "classification languages beyond English, Sinhala, and Singlish.",
        "Passwords for email/password accounts are managed by Firebase Authentication":
            "Passwords for email/password accounts are hashed by the backend and are "
            "never stored as plain text. Google and Facebook credentials remain managed "
            "by their identity providers through Firebase Authentication.",
        "Session tokens and locally cached batch history":
            "Session tokens and the local batch-history cache are stored using encrypted "
            "Android preferences. Authoritative batch-history records are also associated "
            "with the authenticated account on the backend.",
        "The backend records an audit log":
            "The backend stores key account, administration, and message actions in an "
            "append-only SHA-256 hash chain. Verification detects missing, changed, or "
            "out-of-order entries. The design is tamper-evident, not absolutely "
            "tamper-proof; production chain-head anchors should be stored separately in "
            "immutable or append-only storage.",
        "Yes. The system detects English, Sinhala, Singlish, and Tamil.":
            "The active classifier is trained for English, Sinhala, and Singlish. The "
            "Android interface and some report templates can be displayed in Tamil, but "
            "Tamil message classification is outside the active model scope.",
        "The training dataset (English, Sinhala, Singlish, and Tamil)":
            "The active model uses a 10,000-message localized development dataset for "
            "English, Sinhala, and Singlish. Template-generated, localized, and consented "
            "records must remain distinguishable; this development dataset alone does "
            "not establish production accuracy.",
        "The bundled small held-out evaluation set":
            "The bundled 47-message evaluation file is small and includes Tamil messages "
            "outside the active model scope. With the restored model it produces 0.5532 "
            "overall accuracy and must not be treated as the final trilingual benchmark. "
            "A larger independent English/Sinhala/Singlish test set is required.",
        "End-to-end response time depends":
            "End-to-end response time depends on the physical device, deployed backend, "
            "and network. Project benchmark scripts use an in-process FastAPI TestClient "
            "and therefore measure local API integration latency, not real mobile-network "
            "or on-screen response time.",
        "Active model: KRISS-Colab-2.0-Trilingual":
            "Active model: KRISS-Colab-2.0-Trilingual, a complete scikit-learn pipeline "
            "combining word-level and character-level TF-IDF features with a Multinomial "
            "Naive Bayes classifier (alpha = 0.35).",
        "Main training dataset:":
            "Active model dataset: 10,000 localized development messages in English, "
            "Sinhala, and Singlish, spanning the Legitimate, Spam, and Scam classes. "
            "The project must retain accurate provenance for template-generated, "
            "localized, and consented records.",
        "Reported metadata results:":
            "Reported notebook metadata: 0.883 test accuracy and approximately 0.8322 "
            "scam recall. The minimum project threshold is 0.90, while 0.95 is an "
            "aspirational target; the restored model is therefore below the minimum.",
        "Small held-out sample set:":
            "Separate evaluation file: 47 multilingual messages used as a diagnostic "
            "spot-check. It includes Tamil messages outside the active model scope and "
            "produces 0.5532 accuracy with the restored model. It is not a valid final "
            "trilingual accuracy benchmark.",
        "Important: No claim in this manual":
            "Important: Model performance must be reported from measured results. A "
            "credible final claim requires a sufficiently large, independent, preserved "
            "English/Sinhala/Singlish test set with documented source and consent "
            "provenance. The current model does not yet demonstrate the stated 0.90 "
            "minimum on its notebook evaluation.",
        "A.9 Dataset Administration (/api/dataset-admin":
            "A.9 Dataset Administration (/api/admin/dataset, admin role required)",
    }
    for prefix, text in replacements.items():
        # This sentence occurs three times for report-language steps; handle separately.
        if prefix == "Choose the report language: English, Sinhala, or Tamil.":
            continue
        replace_paragraph(document, prefix, text)

    report_paragraphs = [
        p for p in document.paragraphs
        if p.text.strip() in {
            "Choose the report language: English, Sinhala, or Tamil.",
            "Select the report language: English, Sinhala, or Tamil.",
        }
    ]
    if len(report_paragraphs) != 3:
        raise ValueError(f"Expected three report-language paragraphs; found {len(report_paragraphs)}")
    report_paragraphs[0].text = "Choose English, Sinhala, or Tamil as the report language."
    report_paragraphs[1].text = "Choose English, Sinhala, or Tamil as the batch report language."
    report_paragraphs[2].text = "Choose English or Sinhala as the history-item report language."

    # Strengthen operational instructions without claiming evidence that does not exist.
    replace_paragraph(
        document,
        "Tap Create Backup",
        "Tap Create Backup to generate and checksum a database backup on the server. "
        "This manual action is separate from the daily schedule; an administrator must "
        "run backend/scripts/setup_backup_scheduler.ps1 and verify the Windows scheduled task.",
    )
    replace_paragraph(
        document,
        "The Admin Dashboard opens on a summary",
        "The Admin Dashboard presents stored API performance metrics and centralized "
        "error records from the backend. These operational metrics describe observed "
        "requests; they do not by themselves prove a 99% uptime period.",
    )

    # Cover and technology/feature tables.
    cover = document.tables[0]
    cover.cell(2, 1).text = "2.1"
    cover.cell(6, 1).text = "Corrected"

    technology = document.tables[1]
    technology.cell(2, 1).text = (
        "Backend email/password authentication; Firebase Authentication for Google "
        "and Facebook identity federation"
    )
    technology.cell(5, 1).text = (
        "scikit-learn word/character TF-IDF with Multinomial Naive Bayes "
        "(English, Sinhala, Singlish)"
    )

    features = document.tables[2]
    features.cell(7, 1).text = (
        "Interface language (English/Sinhala/Tamil), notification toggle, and automatic "
        "history deletion; Tamil UI localization does not imply Tamil model support"
    )
    features.cell(10, 1).text = (
        "Encrypted local session and batch cache, server-side batch history, rate "
        "limiting, tamper-evident audit logging, session revocation, restricted CORS, "
        "and production HTTPS enforcement"
    )

    glossary = document.tables[5]
    glossary.cell(8, 1).text = (
        "The statistical classifier used by the active "
        "KRISS-Colab-2.0-Trilingual model."
    )

    # Correct API prefixes and make relative-path tables unambiguous.
    for row in document.tables[7].rows[1:]:
        row.cells[1].text = "/api/auth" + row.cells[1].text
    for row in document.tables[8].rows[1:]:
        row.cells[1].text = "/api" + row.cells[1].text
    for row in document.tables[9].rows[1:]:
        suffix = row.cells[1].text
        row.cells[1].text = "/api/history" + suffix
    for row in document.tables[11].rows[1:]:
        row.cells[1].text = "/api" + row.cells[1].text
    for row in document.tables[12].rows[1:]:
        row.cells[1].text = "/api/batch" + row.cells[1].text
    for row in document.tables[13].rows[1:]:
        row.cells[1].text = "/api/admin" + row.cells[1].text
    for row in document.tables[14].rows[1:]:
        row.cells[1].text = "/api/admin/dataset" + row.cells[1].text

    # The community table used placeholders; replace them with the actual paths.
    community_paths = [
        "/api/community-spam",
        "/api/my-reports",
        "/api/reports",
        "/api/my-feedback",
        "/api/feedback",
    ]
    for row, path in zip(document.tables[10].rows[1:], community_paths):
        row.cells[1].text = path

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
